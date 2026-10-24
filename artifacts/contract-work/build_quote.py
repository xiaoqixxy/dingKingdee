from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
out=root/'钉钉金蝶数据同步服务报价单.docx'
d=Document(); s=d.sections[0]
s.page_width=Cm(21); s.page_height=Cm(29.7)
s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.7)
s.left_margin=s.right_margin=Cm(1.9)
for name,size in [('Normal',10.5),('Title',20),('Heading 1',12),('Heading 2',11)]:
    st=d.styles[name];st.font.name='宋体';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.15
for st in d.styles:
    for el in list(st.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
d.core_properties.title='钉钉与金蝶云星空数据同步服务报价单';d.core_properties.author=''
def p(txt,bold=False):
    x=d.add_paragraph();x.add_run(txt).bold=bold
    return x
def h(txt): d.add_paragraph(txt,'Heading 1')
def table(headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Cm(w)
    for c,txt,w in zip(t.rows[0].cells,headers,widths): c.text=txt;c.width=Cm(w)
    for row in rows:
        for c,txt,w in zip(t.add_row().cells,row,widths):c.text=txt;c.width=Cm(w)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if i==0: trpr.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
            tcpr.append(b)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF4');tcpr.append(sh)
            mar=OxmlElement('w:tcMar')
            for side,v in [('top','65'),('bottom','65'),('left','90'),('right','90')]:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),v);el.set(qn('w:type'),'dxa');mar.append(el)
            tcpr.append(mar)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(2)
                if len(headers)>2 and j>1: pp.alignment=1
                for rr in pp.runs:rr.font.size=Pt(9.5);rr.bold=(i==0)
    return t

d.add_paragraph('钉钉与金蝶云星空\n数据同步服务报价单','Title')
p('报价编号：【填写】    报价日期：【年/月/日】    有效期至：【年/月/日】')
p('致：【客户公司全称】\n报价方：【服务提供方全称】\n联系人：【姓名】    电话：【填写】    邮箱：【填写】')
p('本报价用于金蝶云星空数据同步至钉钉 AI 表格的插件订阅及实施服务。采购范围、套餐权益和交付要求以下表及双方确认的配置清单为准。')
h('一  本次采购范围')
p('采购套餐：【套餐名称】    订阅期限：【】个月    钉钉企业：【名称及 CorpId】\n金蝶版本与账套：【填写】    目标钉钉表格：【填写】')
p('套餐权益：绑定表单【】个；单次同步上限【】行；月度额度【无则填无】；同步触发方式及频率【填写】。明细模式按输出明细行计量；其他计量、重试和超限规则须在配置清单确认。')
h('二  费用明细')
table(['项目','计价及范围','数量','含税单价\n人民币元','含税金额\n人民币元'],[
('插件订阅','【套餐】／【计费周期】','【】','【】','【】'),
('首次实施','连接配置、表单与字段核对\n覆盖【】个表单，远程实施','【】次','【】','【】'),
('培训与支持','培训【】次；支持期限【】\n□已含订阅费  □单独收费','【】','【】','【】'),
('选购增项','已确认的增项【编号及名称】\n未选项目不计入总价','【】','【】','【】'),
('费用合计','以上含税金额之和','—','—','【A】'),
('优惠金额','优惠条件：【填写；无则填0】','—','—','【B】'),
('本次报价总额','含税应付总额＝A－B','—','—','【C】'),
],[2.6,6.7,1.6,3.1,3.2])
p('人民币大写：【填写】。币种：人民币；发票类型：【填写】；适用税率：【】%。各项单价和金额按含税口径填写；如税率不同，请逐项注明。不含税金额【】元，税额【】元，价税合计应与总额一致。')

d.add_page_break()
h('三  交付内容与实施安排')
p('标准交付包括订阅开通、约定金蝶访问配置、支持表单及字段核对、约定表头或明细同步设置、筛选与排序设置、操作说明及必要验收测试。表单清单及特殊字段转换要求：【填写或注明附件版本】。')
p('计划工期：【】个工作日，自付款条件满足、甲方提供完整授权与可访问环境且双方确认配置清单之日起计算。客户资料或环境延误时，双方书面调整交付日期。')
p('验收要求：指定账套和表单连接正常；按约定转换规则，样本【】条与源数据一致；明细关联、筛选排序、分页和限额提示符合配置清单。验收期【】个工作日，不符合项由双方记录并确认修复计划。')
h('四  可选增项')
table(['选项','服务及计价范围','计价单位','含税单价'],[
('□','新增表单适配【名称、版本及字段范围】','个表单','【】元'),
('□','特殊字段转换或接口定制【需求编号】','项／人日','【】元'),
('□','额外培训【时长及人数】','场','【】元'),
('□','现场实施【地点及天数；差旅计费方式】','人日','【】元'),
('□','私有化部署【另附部署、许可及运维方案】','套','【】元'),
],[1.3,9.0,2.8,4.1])
p('增项仅在勾选、明确数量与范围并经双方确认后计入第二部分费用。定制功能、私有化部署及其他增项需先确认技术可行性、交付内容、工期和验收标准。')
h('五  商务与服务条款')
p('1. 付款方式：□全额预付  □分期。首付款【】%／【】元，在【】后【】日内支付；尾款【】%／【】元，在【】后【】日内支付。合计应等于本次报价总额。')
p('2. 收款信息：开户名【】；开户行【】；账号【】。开票时间：【填写条件及期限】。线上支付须关联可核验订单，避免重复付款。')
p('3. 支持安排：服务时间【】；渠道【】；响应时限【】。可用性、恢复时限、备份及安全措施，以双方正式合同中的已确认承诺为准。')
p('4. 未包含费用：钉钉与金蝶许可及接口授权、客户网络与硬件费用；未列入明细的历史数据清洗、现场差旅、双向同步、向金蝶写入或审核单据、源码交付。特别包含事项：【无则填无】。')
p('5. 续费与变更：后续续费价格及期限另行确认；升级、降级、增购和超额费用须经客户授权确认。报价有效期内，已确认范围和价格按本报价执行；范围变化需书面重新报价。')
p('6. 本报价供商务确认，正式服务的生效、退款、数据处理、保密与违约责任以双方签署的服务合同为准。客户确认采购时，应同时补齐报价及配置清单中的空白项。')
h('六  报价与采购确认')
p('报价方盖章或授权代表：【】    日期：【】\n客户确认盖章或授权代表：【】    日期：【】\n客户选定套餐及增项编号：【】    确认含税总额：【】元')
f=s.footer.paragraphs[0];f.alignment=2;f.add_run('第 ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld);f.add_run(' 页')
d.save(out); print(out)
